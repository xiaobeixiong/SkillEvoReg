"""Verified-400 metadata fixes and a SpreadsheetBench-compatible evaluator."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

import openpyxl
from openpyxl.utils.cell import range_boundaries

from skillopt.envs.spreadsheetbench.evaluator import _compare_cell_value


# These source records cannot be evaluated as written. Keep corrections explicit
# and auditable instead of silently changing the downloaded dataset.
METADATA_OVERRIDES = {
    "60-7": {
        "answer_position": "'Consolidated Tracker'!A3:E11",
        "answer_sheet": "Consolidated Tracker",
        "reason": "answer_sheet lists source sheets although only Consolidated Tracker is output",
    },
    "283-32": {
        "answer_position": "'Sheet3'!A:G,'Sheet4'!A:G",
        "answer_sheet": "'Sheet3','Sheet4'",
        "reason": "source answer_position has mismatched quotes for two output sheets",
    },
}


def normalize_task_metadata(item: dict) -> dict:
    normalized = deepcopy(item)
    override = METADATA_OVERRIDES.get(str(item["id"]))
    if override:
        normalized.update({key: value for key, value in override.items() if key != "reason"})
    return normalized


def answer_position_for_evaluation(item: dict) -> str:
    normalized = normalize_task_metadata(item)
    answer_position = normalized.get("answer_position", "")
    answer_sheet = normalized.get("answer_sheet", "")
    if answer_position and answer_sheet and "!" not in answer_position:
        return f"{answer_sheet}!{answer_position}"
    return answer_position


def _split_references(answer_position: str) -> list[str]:
    """Split comma-separated Excel references while preserving quoted sheet names."""
    references = []
    current = []
    quote = None
    for index, char in enumerate(answer_position):
        if quote == char:
            quote = None
            current.append(char)
        elif char in ("'", '"') and quote is None:
            # A quote protects commas only in the valid Excel form
            # 'sheet, name'!A1. Stray quotes elsewhere are dataset noise.
            at_reference_start = not "".join(current).strip()
            closing_quote = answer_position.find(char, index + 1)
            next_bang = answer_position.find("!", index + 1)
            if at_reference_start and index < closing_quote < next_bang:
                quote = char
            current.append(char)
        elif char == "," and quote is None:
            reference = "".join(current).strip()
            if reference:
                references.append(reference)
            current = []
        else:
            current.append(char)
    reference = "".join(current).strip()
    if reference:
        references.append(reference)
    return references


def _unquote_sheet_name(sheet_name: str) -> str:
    return sheet_name.strip().strip("'\"").strip("!").replace("''", "'")


def _normalize_cell_range(cell_range: str) -> str:
    cell_range = cell_range.strip().strip("'\"")
    missing_end_column = re.fullmatch(r"([A-Za-z]+)(\d+):(\d+)", cell_range)
    if missing_end_column:
        column, start_row, end_row = missing_end_column.groups()
        return f"{column}{start_row}:{column}{end_row}"
    return cell_range


def compare_workbooks(
    gold_path: str | Path,
    prediction_path: str | Path,
    answer_position: str,
    *,
    formula_sensitive: bool = False,
) -> tuple[bool, str]:
    """Compare answer cells with official value semantics and robust range parsing."""
    if not Path(prediction_path).exists():
        return False, "file not exist"
    try:
        gold = openpyxl.load_workbook(gold_path, data_only=not formula_sensitive)
        prediction = openpyxl.load_workbook(
            prediction_path, data_only=not formula_sensitive
        )
    except Exception as exc:  # noqa: BLE001
        return False, f"load error: {exc}"

    try:
        references = _split_references(answer_position or "")
        if not references:
            return False, "empty answer position"
        for reference in references:
            if "!" in reference:
                sheet_text, cell_range = reference.rsplit("!", 1)
                sheet_name = _unquote_sheet_name(sheet_text)
            else:
                sheet_name = gold.sheetnames[0]
                cell_range = reference
            if sheet_name not in gold.sheetnames:
                return False, f"gold worksheet not found: {sheet_name}"
            if sheet_name not in prediction.sheetnames:
                return False, f"prediction worksheet not found: {sheet_name}"

            gold_sheet = gold[sheet_name]
            prediction_sheet = prediction[sheet_name]
            cell_range = _normalize_cell_range(cell_range)
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
            min_col = min_col or 1
            min_row = min_row or 1
            max_col = max_col or max(gold_sheet.max_column, prediction_sheet.max_column)
            max_row = max_row or max(gold_sheet.max_row, prediction_sheet.max_row)
            for row in range(min_row, max_row + 1):
                for column in range(min_col, max_col + 1):
                    gold_value = gold_sheet.cell(row=row, column=column).value
                    prediction_value = prediction_sheet.cell(row=row, column=column).value
                    if formula_sensitive and gold_sheet.cell(row=row, column=column).data_type == "f":
                        if not _compare_formula(gold_value, prediction_value):
                            coordinate = gold_sheet.cell(row=row, column=column).coordinate
                            return (
                                False,
                                f"formula@{sheet_name}!{coordinate}: "
                                f"gt={gold_value!r} pred={prediction_value!r}",
                            )
                    elif not _compare_cell_value(gold_value, prediction_value):
                        coordinate = gold_sheet.cell(row=row, column=column).coordinate
                        return (
                            False,
                            f"value@{sheet_name}!{coordinate}: "
                            f"gt={gold_value!r} pred={prediction_value!r}",
                        )
        return True, ""
    except Exception as exc:  # noqa: BLE001
        return False, f"{type(exc).__name__}: {exc}"
    finally:
        gold.close()
        prediction.close()


def _compare_formula(gold_value: object, prediction_value: object) -> bool:
    """Compare simple generated formulas while allowing commutative products."""
    if not isinstance(gold_value, str) or not gold_value.startswith("="):
        return False
    if not isinstance(prediction_value, str) or not prediction_value.startswith("="):
        return False

    def normalize(value: str) -> str:
        return re.sub(r"\s+|\$", "", value).upper()

    gold_formula = normalize(gold_value)
    prediction_formula = normalize(prediction_value)
    if gold_formula == prediction_formula:
        return True

    def product_references(value: str) -> tuple[str, str] | None:
        binary = re.fullmatch(r"=([A-Z]+\d+)\*([A-Z]+\d+)", value)
        if binary:
            return tuple(sorted(binary.groups()))
        function = re.fullmatch(
            r"=PRODUCT\(([A-Z]+\d+),([A-Z]+\d+)\)", value
        )
        if function:
            return tuple(sorted(function.groups()))
        return None

    gold_product = product_references(gold_formula)
    return gold_product is not None and product_references(prediction_formula) == gold_product


def evaluate(
    prediction_path: str | Path,
    gold_path: str | Path,
    instruction_type: str,
    answer_position: str,
) -> dict:
    ok, reason = compare_workbooks(
        gold_path,
        prediction_path,
        answer_position,
        formula_sensitive=instruction_type == "Formula-Level Manipulation",
    )
    return {"ok": ok, "reason": reason, "instruction_type": instruction_type}


def verification_report(
    prediction_path: str | Path,
    gold_path: str | Path,
    answer_position: str,
) -> str:
    """Compact replacement for SkillOpt's parser-sensitive feedback report."""
    ok, reason = compare_workbooks(gold_path, prediction_path, answer_position)
    if ok:
        return "## Output Verification\nAll evaluated answer cells match."
    return f"## Output Verification\nFirst mismatch: {reason}"
